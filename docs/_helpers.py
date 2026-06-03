"""
SendFloo — Helpers communs pour générer les documents Word.

Toutes les fonctions ici sont utilisées par les générateurs de chaque catégorie
de document (cadrage, métier, technique, etc.).
"""
from __future__ import annotations
import os
from datetime import date
from pathlib import Path
from typing import List, Tuple

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# Palette SendFloo
COLOR_NAVY = RGBColor(0x02, 0x2A, 0x6B)
COLOR_NAVY_DARK = RGBColor(0x01, 0x16, 0x45)
COLOR_EMERALD = RGBColor(0x04, 0xD4, 0x6F)
COLOR_ORANGE = RGBColor(0xFF, 0xA5, 0x00)
COLOR_PURPLE = RGBColor(0x7C, 0x3A, 0xED)
COLOR_GOLD = RGBColor(0xC9, 0xA2, 0x27)
COLOR_GREEN = RGBColor(0x04, 0x78, 0x57)
COLOR_GREY_DARK = RGBColor(0x36, 0x3F, 0x5C)
COLOR_GREY_MID = RGBColor(0x6E, 0x76, 0x9A)
COLOR_GREY_LIGHT = RGBColor(0xE5, 0xE9, 0xF2)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_BLACK = RGBColor(0x0F, 0x1F, 0x4E)

OUTPUT_DIR = Path("/app/docs/sendfloo_documentation")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ========================================================================
# Styles de base
# ========================================================================

def init_doc() -> Document:
    """Crée un Document avec marges raisonnables et fonts par défaut."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
    # Style global
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    return doc


def _set_cell_bg(cell, hex_color: str):
    """Définit la couleur de fond d'une cellule de tableau (hex sans #)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


# ========================================================================
# Cover page
# ========================================================================

def add_cover(doc: Document, *, title: str, subtitle: str, doc_code: str,
              version: str = "1.0", auteurs: str = "Direction Produit SendFloo",
              destinataires: str = "Comité de Pilotage, Équipes Technique & Métier"):
    """Ajoute une page de couverture stylisée."""
    # Marque
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("SENDFLOO ")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = COLOR_NAVY
    r2 = p.add_run("· documentation projet")
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_GREY_MID

    # Espace
    for _ in range(6):
        doc.add_paragraph()

    # Titre
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = COLOR_NAVY_DARK

    # Sous-titre
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(subtitle)
    r.font.size = Pt(16)
    r.font.color.rgb = COLOR_GREY_DARK

    for _ in range(10):
        doc.add_paragraph()

    # Tableau métadonnées
    tbl = doc.add_table(rows=5, cols=2)
    tbl.style = "Light Grid Accent 1"
    rows = [
        ("Code document", doc_code),
        ("Version", version),
        ("Date d'édition", date.today().strftime("%d/%m/%Y")),
        ("Auteur(s)", auteurs),
        ("Destinataires", destinataires),
    ]
    for i, (label, value) in enumerate(rows):
        c0, c1 = tbl.rows[i].cells
        c0.text = label
        c1.text = value
        for p in c0.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = COLOR_NAVY
                r.font.size = Pt(11)

    doc.add_page_break()


# ========================================================================
# Sommaire (TOC manuel pour fiabilité)
# ========================================================================

def add_toc(doc: Document, entries: List[Tuple[str, int]]):
    """
    Ajoute un sommaire manuel.
    entries: [(titre, niveau)], niveau 1 = section principale, 2 = sous-section.
    """
    h = doc.add_heading("Sommaire", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        run.font.color.rgb = COLOR_NAVY_DARK
    for title, lvl in entries:
        p = doc.add_paragraph()
        indent = "    " * (lvl - 1)
        r = p.add_run(f"{indent}{title}")
        r.font.size = Pt(11 if lvl == 1 else 10)
        if lvl == 1:
            r.bold = True
            r.font.color.rgb = COLOR_NAVY_DARK
        else:
            r.font.color.rgb = COLOR_GREY_DARK
    doc.add_page_break()


# ========================================================================
# Sections / Sous-sections / Paragraphes
# ========================================================================

def h1(doc: Document, text: str):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = COLOR_NAVY_DARK
        run.font.size = Pt(20)
    return h


def h2(doc: Document, text: str):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = COLOR_NAVY
        run.font.size = Pt(15)
    return h


def h3(doc: Document, text: str):
    h = doc.add_heading(text, level=3)
    for run in h.runs:
        run.font.color.rgb = COLOR_GREY_DARK
        run.font.size = Pt(13)
    return h


def p(doc: Document, text: str, *, italic: bool = False, bold: bool = False, color=None):
    par = doc.add_paragraph()
    r = par.add_run(text)
    r.italic = italic
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return par


def bullet(doc: Document, items: List[str]):
    for item in items:
        par = doc.add_paragraph(item, style="List Bullet")


def numbered(doc: Document, items: List[str]):
    for item in items:
        par = doc.add_paragraph(item, style="List Number")


def table(doc: Document, headers: List[str], rows: List[List[str]], *, header_color: str = "022A6B"):
    """Ajoute un tableau stylisé."""
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Light Grid Accent 1"
    # Header
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = h
        _set_cell_bg(cell, header_color)
        for par in cell.paragraphs:
            for r in par.runs:
                r.bold = True
                r.font.color.rgb = COLOR_WHITE
                r.font.size = Pt(10)
    # Rows
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.rows[i + 1].cells[j]
            cell.text = str(val)
            for par in cell.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(10)
    return tbl


def callout(doc: Document, title: str, content: str, *, color="FFA500"):
    """Encart d'attention (couleur orange par défaut)."""
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    _set_cell_bg(cell, color)
    par = cell.paragraphs[0]
    r = par.add_run(f"⚡ {title}\n")
    r.bold = True
    r.font.color.rgb = COLOR_WHITE
    r.font.size = Pt(12)
    r2 = par.add_run(content)
    r2.font.color.rgb = COLOR_WHITE
    r2.font.size = Pt(10)
    doc.add_paragraph()


def page_break(doc: Document):
    doc.add_page_break()


# ========================================================================
# Save helper
# ========================================================================

def save(doc: Document, filename: str) -> Path:
    """Sauvegarde le document et retourne le chemin."""
    path = OUTPUT_DIR / filename
    doc.save(str(path))
    print(f"  ✅ {filename} ({path.stat().st_size // 1024} KB)")
    return path
