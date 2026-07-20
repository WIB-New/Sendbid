#!/usr/bin/env python3
"""Génère un audit détaillé du projet SendBID / PayBID au format Word (.docx)."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import datetime


def set_heading_style(doc, level, font_size, color):
    """Configure un style de titre personnalisé."""
    style = doc.styles[f"Heading {level}"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(font_size)
    font.bold = True
    font.color.rgb = RGBColor(*color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_paragraph(doc, text, bold=False, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def add_bullet(doc, text, indent_level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + indent_level * 0.25)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.name = "Calibri"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    return table


def main():
    doc = Document()

    # Styles
    set_heading_style(doc, 1, 18, (2, 42, 94))
    set_heading_style(doc, 2, 14, (2, 84, 147))
    set_heading_style(doc, 3, 12, (68, 114, 196))

    # Page de titre
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Audit fonctionnel détaillé")
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(2, 42, 94)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Projet SendBID / PayBID")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(2, 84, 147)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run(f"Généré le {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}")
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    doc.add_page_break()

    # Résumé exécutif
    add_heading(doc, "Résumé exécutif", level=1)
    add_paragraph(doc,
        "Ce document présente l'état détaillé des fonctionnalités déjà implémentées et opérationnelles "
        "dans le projet SendBID (application client) et PayBID (application agent). Le backend est "
        "développé en FastAPI avec une base de données MongoDB.")
    add_bullet(doc, "SendBID : application client pour envoyer/recevoir de l'argent, recharger et retirer.")
    add_bullet(doc, "PayBID : application agent pour les opérations de caisse (cash-in, cash-out, scan QR).")
    add_bullet(doc, "Backend FastAPI : authentification, wallet, paiements externes, admin panel.")
    add_bullet(doc, "Prérequis externes : clés API Stripe, PayPal, Twilio et SendGrid nécessaires pour les paiements réels.")

    # Partie 1 : SendBID
    add_heading(doc, "1. SendBID — Application client", level=1)

    add_heading(doc, "1.1 Authentification & onboarding", level=2)
    add_paragraph(doc, "Le flux d'inscription et de connexion est fonctionnel.", bold=True)
    add_table(doc,
        ["Fonctionnalité", "État", "Fichiers / Détails"],
        [
            ["Inscription email + téléphone + pays", "Opérationnel", "sendbid/app/(auth)/signup.tsx + backend/routers/auth/core.py::register"],
            ["Validation OTP email + téléphone", "Opérationnel", "Codes générés côté backend, envoyés par Twilio/SendGrid, visibles en dev"],
            ["Login email / téléphone / profile_id", "Opérationnel", "backend/routers/auth/core.py::login"],
            ["Création de PIN à 6 chiffres", "Opérationnel", "Obligatoire après inscription, gating via sendbid/app/_layout.tsx"],
            ["Authentification biométrique", "Partiel", "Champ présent côté backend, activation UI non auditée"],
        ]
    )

    add_heading(doc, "1.2 Portefeuille (Wallet)", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Détails"],
        [
            ["Affichage du solde + devise", "Opérationnel", "Basé sur user.country et wallet.currency"],
            ["Historique des transactions", "Opérationnel", "Onglet wallet"],
            ["Sessions actives", "Opérationnel", "backend/routers/sessions.py"],
        ]
    )

    add_heading(doc, "1.3 Recharge du compte", level=2)
    add_paragraph(doc,
        "La recharge peut se faire par carte bancaire (Stripe), PayPal, portefeuille mobile ou espèces. "
        "Les recharges par Mobile Money et PayPal forcent désormais la sélection d'un compte lié.")
    add_table(doc,
        ["Méthode", "État", "Détails"],
        [
            ["Carte bancaire (Stripe)", "Connecté, clé API requise", "backend/services/payments.py utilise STRIPE_API_KEY"],
            ["PayPal", "Connecté, clés API requises", "backend/routers/paypal.py utilise PAYPAL_CLIENT_ID + PAYPAL_SECRET"],
            ["Portefeuille mobile (MoMo)", "Opérationnel", "Forçage de la sélection d'un compte lié dans recharge.tsx"],
            ["Espèces", "Opérationnel", "Génération de QR code à présenter à un agent PayBID"],
        ]
    )
    add_paragraph(doc, "Améliorations récentes :")
    add_bullet(doc, "Le libellé 'Mobile Money' a été renommé en 'Portefeuille mobile'.")
    add_bullet(doc, "Si aucun compte lié n'existe, l'utilisateur est redirigé vers Profil → Comptes liés.")
    add_bullet(doc, "Les packages de recharge sont convertis dans la devise locale du pays.")

    add_heading(doc, "1.4 Retrait d'argent", level=2)
    add_paragraph(doc,
        "Le retrait est possible en espèces (chez un agent), par virement bancaire, "
        "par portefeuille mobile ou par PayPal. La sélection d'un compte lié est obligatoire.")
    add_table(doc,
        ["Méthode", "État", "Détails"],
        [
            ["Espèces (QR agent)", "Opérationnel", "backend/routers/wallet/withdraw.py génère un QR"],
            ["Virement bancaire", "Opérationnel", "Nécessite un compte bancaire lié actif"],
            ["Portefeuille mobile", "Opérationnel", "Nécessite un compte MoMo lié actif"],
            ["PayPal", "Opérationnel", "Nécessite un compte PayPal lié actif"],
        ]
    )
    add_paragraph(doc, "Améliorations récentes :")
    add_bullet(doc, "Sélection obligatoire d'un compte lié pour banque / MoMo / PayPal.")
    add_bullet(doc, "Pré-sélection automatique du premier compte actif correspondant.")
    add_bullet(doc, "Suppression de la saisie manuelle pour bank / MoMo / PayPal.")
    add_bullet(doc, "Tous les pays disposent désormais des 4 méthodes de retrait.")

    add_heading(doc, "1.5 Comptes liés", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Fichiers"],
        [
            ["CRUD comptes bancaires / MoMo / PayPal", "Opérationnel", "sendbid/app/wallet/linked-accounts.tsx"],
            ["API backend dédiée", "Opérationnel", "backend/routers/wallet/linked_accounts.py"],
            ["Validation automatique (format)", "Opérationnel", "Endpoint /linked-accounts/{id}/verify"],
            ["Liste de banques par pays", "Opérationnel", "sendbid/src/currency.ts"],
            ["Liste d'opérateurs MoMo par pays", "Opérationnel", "sendbid/src/currency.ts"],
            ["Évitement des doublons", "Opérationnel", "Vérifié côté backend"],
        ]
    )

    add_heading(doc, "1.6 Transferts et bénéficiaires", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Détails"],
        [
            ["Envoi à un autre utilisateur SendBID", "Opérationnel", "sendbid/app/transfer/new.tsx"],
            ["Liste de bénéficiaires", "Opérationnel", "sendbid/app/beneficiaries.tsx"],
            ["Envoi à un agent (cash-out)", "Opérationnel", "Via retrait espèces"],
        ]
    )

    add_heading(doc, "1.7 Profil & paramètres", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Détails"],
        [
            ["Modification infos personnelles", "Opérationnel", "sendbid/app/personal-info.tsx"],
            ["Internationalisation", "Opérationnel", "FR, EN, ES, DE, IT, PT, AR dans sendbid/src/i18n/"],
            ["Notifications push", "Présent", "Gestion dans _layout.tsx"],
        ]
    )

    add_heading(doc, "1.8 Gestion des erreurs API", level=2)
    add_paragraph(doc,
        "La fonction apiError() a été améliorée dans sendbid/src/api.ts et sendbid/src/src/api.ts. "
        "Elle privilégie le message exact du backend et fournit des fallbacks par code HTTP.")
    add_bullet(doc, "Messages explicites pour PIN incorrect, session invalide, accès refusé, etc.")
    add_bullet(doc, "L'erreur générique 'Email ou mot de passe incorrect' ne masque plus les erreurs backend.")

    # Partie 2 : PayBID
    doc.add_page_break()
    add_heading(doc, "2. PayBID — Application agent", level=1)
    add_paragraph(doc,
        "PayBID est la variante agent de la même codebase Expo. Elle permet aux agents agréés "
        "de réaliser des opérations de caisse.")

    add_heading(doc, "2.1 Authentification agent", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Détails"],
        [
            ["Login agent", "Opérationnel", "sendbid/app/paybid/login.tsx"],
            ["Signup agent", "Opérationnel", "sendbid/app/paybid/signup.tsx"],
            ["PIN agent", "Opérationnel", "Même mécanisme que le client"],
            ["Variante paybid", "Opérationnel", "EXPO_PUBLIC_APP_VARIANT=paybid"],
        ]
    )

    add_heading(doc, "2.2 Tableau de bord agent", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Détails"],
        [
            ["Solde de caisse (float)", "Opérationnel", "Onglet Float"],
            ["Historique des mouvements", "Opérationnel", "sendbid/app/paybid/movements.tsx"],
            ["Statistiques agent", "Opérationnel", "sendbid/app/paybid/(tabs)/account.tsx"],
        ]
    )

    add_heading(doc, "2.3 Opérations de caisse", level=2)
    add_table(doc,
        ["Fonctionnalité", "État", "Détails"],
        [
            ["Cash-in (dépôt espèces)", "Opérationnel", "sendbid/app/paybid/cash-recharge.tsx"],
            ["Cash-out (retrait espèces)", "Opérationnel", "Scan du QR de retrait client"],
            ["Scan QR", "Opérationnel", "sendbid/app/paybid/scan.tsx"],
            ["Gestion de la caisse (float)", "Opérationnel", "sendbid/app/paybid/float.tsx"],
            ["Opérations d'agence", "Opérationnel", "sendbid/app/paybid/agency-ops.tsx"],
        ]
    )

    add_heading(doc, "2.4 Transferts inter-agents", level=2)
    add_bullet(doc, "Transferts entre agents opérationnels via sendbid/app/paybid/transfer/")

    # Partie 3 : Backend
    doc.add_page_break()
    add_heading(doc, "3. Backend FastAPI", level=1)

    add_heading(doc, "3.1 Authentification", level=2)
    add_table(doc,
        ["Module", "État", "Détails"],
        [
            ["Register / Login / Me / UpdateMe", "Opérationnel", "backend/routers/auth/core.py"],
            ["PIN hash + vérification", "Opérationnel", "backend/core/security.py + require_pin"],
            ["JWT Bearer token", "Opérationnel", "backend/core/deps.py"],
            ["OTP email/téléphone", "Opérationnel", "backend/routers/auth/otp.py + services/notify.py"],
            ["Sessions actives", "Opérationnel", "backend/routers/sessions.py"],
        ]
    )

    add_heading(doc, "3.2 Wallet", level=2)
    add_table(doc,
        ["Module", "État", "Détails"],
        [
            ["Solde et transactions", "Opérationnel", "backend/routers/wallet/"],
            ["Retrait QR (cash-out agent)", "Opérationnel", "backend/routers/wallet/withdraw.py"],
            ["Recharge espèces", "Opérationnel", "backend/routers/wallet/recharge.py"],
            ["Comptes liés", "Opérationnel", "backend/routers/wallet/linked_accounts.py"],
        ]
    )

    add_heading(doc, "3.3 Paiements externes", level=2)
    add_table(doc,
        ["Module", "État", "Détails"],
        [
            ["Stripe", "Intégré, clé requise", "backend/services/payments.py — checkout + webhook"],
            ["PayPal", "Intégré, clés requises", "backend/routers/paypal.py — order + capture"],
            ["Conversion de devises", "Opérationnel", "Packages convertis dans la devise locale"],
        ]
    )

    add_heading(doc, "3.4 Admin / Monitoring", level=2)
    add_bullet(doc, "Panel web admin : backend/routers/web_panels/admin.py + backend/templates/panels/admin.html")
    add_bullet(doc, "Visualisation des transactions, utilisateurs, KYC depuis MongoDB")

    # Partie 4 : Pays et devises
    doc.add_page_break()
    add_heading(doc, "4. Devises et méthodes par pays", level=1)
    add_paragraph(doc,
        "Tous les pays listés disposent désormais des 4 méthodes : espèces, banque, portefeuille mobile et PayPal.")
    add_table(doc,
        ["Pays", "Devise", "MoMo", "Banque", "PayPal", "Cash"],
        [
            ["Cameroun", "XAF", "Oui", "Oui", "Oui", "Oui"],
            ["Sénégal", "XOF", "Oui", "Oui", "Oui", "Oui"],
            ["Côte d'Ivoire", "XOF", "Oui", "Oui", "Oui", "Oui"],
            ["France", "EUR", "Lydia / PayLib", "Oui", "Oui", "Oui"],
            ["Belgique", "EUR", "Payconiq", "Oui", "Oui", "Oui"],
            ["Allemagne", "EUR", "N26 / Klarna", "Oui", "Oui", "Oui"],
            ["Espagne", "EUR", "Bizum", "Oui", "Oui", "Oui"],
            ["Italie", "EUR", "Satispay", "Oui", "Oui", "Oui"],
            ["Suisse", "CHF", "TWINT", "Oui", "Oui", "Oui"],
            ["Royaume-Uni", "GBP", "Monzo / Revolut", "Oui", "Oui", "Oui"],
            ["États-Unis", "USD", "Venmo / Zelle / Cash App", "Oui", "Oui", "Oui"],
            ["Canada", "CAD", "Interac e-Transfer", "Oui", "Oui", "Oui"],
            ["Maroc", "MAD", "Oui", "Oui", "Oui", "Oui"],
            ["Tunisie", "TND", "Oui", "Oui", "Oui", "Oui"],
            ["Algérie", "DZD", "Oui", "Oui", "Oui", "Oui"],
            ["Congo, Gabon, Tchad, Mali, Burkina Faso, Togo, Bénin, Niger, RDC, Guinée, Madagascar", "—", "Oui", "Oui", "Oui", "Oui"],
        ]
    )

    # Partie 5 : APIs externes
    doc.add_page_break()
    add_heading(doc, "5. État des APIs externes", level=1)
    add_table(doc,
        ["API", "Intégration", "Configuration requise", "État sans clé"],
        [
            ["Stripe", "Checkout + PaymentSheet", "STRIPE_API_KEY", "'STRIPE_API_KEY is not configured'"],
            ["PayPal", "Orders v2 + Payouts", "PAYPAL_CLIENT_ID, PAYPAL_SECRET, PAYPAL_MODE", "'PayPal non configuré'"],
            ["Twilio", "Validation téléphone + SMS OTP", "Clés Twilio", "Lookup/SMS désactivé en fallback"],
            ["SendGrid", "Email OTP", "Clé SendGrid", "Email OTP simulé en dev"],
        ]
    )
    add_paragraph(doc,
        "Conclusion : les routes et la logique sont connectées, mais elles nécessitent des clés d'API valides "
        "dans les variables d'environnement du backend pour fonctionner en production ou en test.")

    # Partie 6 : Fichiers modifiés
    doc.add_page_break()
    add_heading(doc, "6. Fichiers créés ou modifiés récemment", level=1)

    add_heading(doc, "6.1 Fichiers créés", level=2)
    add_bullet(doc, "sendbid/app/wallet/linked-accounts.tsx — gestion des comptes liés (client).")
    add_bullet(doc, "backend/routers/wallet/linked_accounts.py — API backend des comptes liés.")
    add_bullet(doc, "sendbid/services/ — services auxiliaires.")

    add_heading(doc, "6.2 Fichiers modifiés", level=2)
    add_bullet(doc, "sendbid/src/api.ts + sendbid/src/src/api.ts — meilleure gestion des erreurs API.")
    add_bullet(doc, "sendbid/app/wallet/recharge.tsx — forçage compte lié pour MoMo/PayPal.")
    add_bullet(doc, "sendbid/app/wallet/withdraw.tsx — forçage compte lié pour banque/MoMo/PayPal.")
    add_bullet(doc, "sendbid/src/currency.ts — activation de toutes les méthodes pour tous les pays.")
    add_bullet(doc, "backend/routers/auth/core.py — registration OTP, has_pin, normalisation pays.")
    add_bullet(doc, "sendbid/src/i18n/*.ts — nouvelles traductions.")
    add_bullet(doc, "sendbid/app/beneficiaries.tsx, sendbid/app/transfer/new.tsx, sendbid/app/personal-info.tsx — corrections mineures.")
    add_bullet(doc, "backend/routers/payments.py, backend/routers/wallet/__init__.py — intégration packages et comptes liés.")
    add_bullet(doc, "backend/routers/web_panels/admin.py + backend/templates/panels/admin.html — panel admin.")

    # Partie 7 : Tests et vigilance
    doc.add_page_break()
    add_heading(doc, "7. Points de vigilance et tests recommandés", level=1)
    add_bullet(doc, "Vérifier les clés API Stripe / PayPal / Twilio / SendGrid dans le .env backend.")
    add_bullet(doc, "Tester un retrait avec PIN correct (demo@sendbid.com / PIN 123456) pour valider le message d'erreur.")
    add_bullet(doc, "Tester un retrait avec un compte lié pour chaque méthode (bank, momo, paypal).")
    add_bullet(doc, "Vérifier l'affichage des opérateurs MoMo en France / Canada / USA.")
    add_bullet(doc, "Compiler TypeScript : tsc --noEmit passe actuellement.")
    add_bullet(doc, "Tester l'ajout d'un compte lié et sa vérification automatique.")
    add_bullet(doc, "Vérifier que l'URL /admin du panel admin est accessible et protégée.")

    # Sauvegarde
    output_path = r"c:\Users\Utilisateur\Desktop\Sendfloo.SendBID-06juin2026\AUDIT_SendBID_PayBID.docx"
    doc.save(output_path)
    print(f"Fichier Word généré : {output_path}")


if __name__ == "__main__":
    main()
